# -*- coding: utf-8 -*-
"""팀 공유용 구글 시트 대시보드 (xlsx → 구글 시트 변환 업로드).

일별 실적은 '일별 기록' 탭에 팀원이 직접 입력. 노란 칸만 채우면 나머지는 수식.
"""
import datetime, unicodedata as ud
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DAYS, START = 7, datetime.date(2026, 9, 22)
CAMPS = [("브랜드 방어","24281447473","브랜드_전체",10000,1500,17),
         ("붓기","24270384426","붓기_일반",20000,2500,14),
         ("칼륨","24275789225","칼륨_일반",10000,2000,11)]
GA4 = [("2026-09-09",1007,19,774700),("2026-09-10",1313,13,585300),("2026-09-11",1178,12,505340),
 ("2026-09-12",1597,24,2094824),("2026-09-13",932,15,1202910),("2026-09-14",934,11,1282462),
 ("2026-09-15",1304,25,2984991),("2026-09-16",1193,14,1738612),("2026-09-17",1246,11,1278524),
 ("2026-09-18",1094,11,973840),("2026-09-19",964,16,2342940),("2026-09-20",1251,19,2062900),
 ("2026-09-21",753,5,182560),("2026-09-22",469,9,895220)]
BASE_N = 12
KW = {
 "브랜드 방어": [("젠제네틱스","완전"),("젠제네틱스","구문"),("젠제네틱스 칼륨","구문"),
  ("젠제네틱스 붓기","구문"),("젠제네틱스 후기","구문"),("젠제네틱스 공식몰","구문"),
  ("젠제네틱스 마그네슘","구문"),("젠제네틱스 비타민B","구문"),("젠제네틱스 붓기파우더","구문"),
  ("zengenetics","완전"),("zengenetics","구문"),("zen genetics","구문"),("zengenetic","구문"),
  ("zengenetics korea","구문"),("zengenetics official","구문"),
  ("위스팟바이오랩","구문"),("wespot biolab","구문")],
 "붓기": [("붓기 영양제","구문"),("붓기 제거","구문"),("붓기 빼는 법","구문"),("붓기파우더","구문"),
  ("붓기 보조제","구문"),("얼굴 붓기","구문"),("다리 붓기","구문"),("종아리 붓기","구문"),
  ("아침 붓기","구문"),("붓기 차","구문"),("부종 영양제","구문"),("수분 배출","구문"),
  ("붓기 영양제","확장"),("붓기 칼륨","확장")],
 "칼륨": [("칼륨 영양제","구문"),("칼륨 보충제","구문"),("포타슘 영양제","구문"),("구연산 칼륨","구문"),
  ("칼륨 파우더","구문"),("칼륨 스틱","구문"),("potassium 영양제","구문"),("칼륨 효능","구문"),
  ("칼륨 부족","구문"),("전해질 영양제","구문"),("칼륨 영양제","확장")]}
NEG = [("의약품·의료",["의약품","이뇨제","다이어트약","병원","처방","약국"]),
 ("무료·체험",["무료","공짜"]),("부정 이슈",["부작용","후유증","소송","리콜"]),
 ("기존 고객 CS",["환불","해지"]),("구직",["알바","채용","구인","연봉"]),
 ("투자",["주가","상장"]),("B2B",["도매","총판","대리점","제조사","OEM","ODM"]),
 ("직접 만들기",["레시피","만들기","DIY"]),("무관",["토렌트","중고"])]
HL = {
"브랜드 방어": ["젠제네틱스 공식몰","젠제네틱스 붓기파우더","붓기파우더 공식 판매처","젠제네틱스 칼륨 525mg",
 "올리브영 건강식품 1위","쿠팡 붓기 검색 1위","후기 333건 평점 4.77","재구매율 92%",
 "설인아 PICK 젠제네틱스","물없이 3초 분말스틱","독일 구연산 칼륨 원료","공식몰 최저가 구매",
 "젠제네틱스 Zengenetics 공식몰","무료배송 당일출고","생애 첫 구매 특가"],
"붓기": ["붓기파우더 젠제네틱스","붓기 영양제 칼륨","아침 얼굴 저녁 종아리","짠 음식 즐기는 분께",
 "칼륨 525mg 1일 2포","물없이 먹는 분말스틱","쿠팡 붓기 검색 1위","올리브영 건강식품 1위",
 "후기 333건 4.77점","독일 구연산 칼륨","칼륨 단일 설계","하루 2포 간편 루틴",
 "공식몰에서 구매","첫 구매 특가 확인","무료배송"],
"칼륨": ["칼륨 영양제 525mg","구연산 칼륨 1포 262mg","칼륨 1일기준치 15%","독일 Jungbunzlauer",
 "칼륨 단일 원료 설계","분말스틱 물없이 섭취","젠제네틱스 포타슘","칼륨 보충 하루 2포",
 "후기 333건 평점 4.77","올리브영 건강식품 1위","쿠팡 붓기 검색 1위","EP USP FCC 기준 충족",
 "공식몰 구매하기","무료배송 당일출고","첫 구매 특가"]}
DS = {
"브랜드 방어": ["젠제네틱스 공식몰입니다. 칼륨·마그네슘·비타민B 정품만 판매합니다.",
 "구매 후기 333건 평균 4.77점. 재구매율 92%. 공식몰에서 확인하세요.",
 "올리브영 건강식품 실시간 랭킹 1위 기록. 무료배송·당일출고.",
 "비공식 판매처 구입 시 품질 보상이 불가합니다. 공식몰에서 구매하세요."],
"붓기": ["칼륨 525mg 하루 2포. 물 없이 먹는 분말스틱으로 간편하게 챙기세요.",
 "쿠팡 붓기 검색 1위·올리브영 건강식품 1위. 후기 333건 4.77점.",
 "독일 Jungbunzlauer 구연산 칼륨 단일 설계. 1포당 칼륨 262.5mg.",
 "생애 첫 구매 특가 진행 중. 데이팩 6포로 가볍게 시작해보세요."],
"칼륨": ["1포당 구연산 칼륨 262.5mg, 2포 525mg으로 1일 기준치의 15%입니다.",
 "독일 Jungbunzlauer 원료. EP·USP·FCC·EU 231/2012 기준 충족.",
 "정제가 아닌 분말스틱. 물 없이 뜯어서 바로 섭취할 수 있습니다.",
 "구매 후기 333건 평균 4.77점. 공식몰에서 첫 구매 특가 확인하세요."]}
TODO = [("위험","스니펫 애셋 423881840386 삭제",
  "값이 '비타민B오플렉스'로 잘못 생성. 브랜드 방어 캠페인 → 애셋에서 삭제. API로는 불가. 올바른 것(423782247365)은 등록됨"),
 ("주의","확장검색 3건 감시",
  "붓기 영양제 / 붓기 칼륨 / 칼륨 영양제. 검색어 보고서 보고 엉뚱한 검색은 제외 키워드로 막기"),
 ("주의","심의 비승인 가능 문안",
  "설인아 PICK(유명인 이름) / 올리브영·쿠팡 1위(최상급). 문안 단위 심의라 광고 전체는 계속 게재"),
 ("정리","중복 광고 2개 삭제(선택)",
  "199923244786~825567922000 / 207170880944~825525918138. 이미 일시중지, 지워도 무해")]

NAVY, ACC, IN_F, CA_F, BAND = "12343C", "0F6E6E", "FFF6E3", "EDF3F4", "DCE7E8"
HF   = Font(color="FFFFFF", bold=True, size=10)
TI   = Font(bold=True, size=16, color=NAVY)
SUB  = Font(size=9, color="7A8A90")
SEC  = Font(bold=True, size=11, color=ACC)
BD   = Font(bold=True)
RED  = Font(bold=True, color="A82318")
ORG  = Font(bold=True, color="9A5B00")
GRN  = Font(bold=True, color="1B7F4B")
BIG  = Font(bold=True, size=20, color=NAVY)
CEN  = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(vertical="top", wrap_text=True)
WON, PCT, NUM = '#,##0"원"', '0.00%', '#,##0'
DOW = ["월","화","수","목","금","토","일"]
def glen(s): return sum(2 if ud.east_asian_width(c) in ("W","F") else 1 for c in str(s))
def fill(c): return PatternFill("solid", fgColor=c)

def head(ws, row, cols, widths):
    for i,(c,w) in enumerate(zip(cols,widths),1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill = fill(NAVY); cell.font = HF; cell.alignment = CEN
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 26

wb = Workbook()
avg_t = sum(t for _,_,t,_ in GA4[:BASE_N])/BASE_N
avg_r = sum(v for _,_,_,v in GA4[:BASE_N])/BASE_N
avg_s = sum(s for _,s,_,_ in GA4[:BASE_N])/BASE_N

# ═══ 1. 대시보드 ═══
ws = wb.active; ws.title = "대시보드"
ws.sheet_view.showGridLines = False
ws["A1"] = "젠제네틱스 구글 검색광고"; ws["A1"].font = TI
ws["A2"] = "계정 798-667-9268 · 통화 KRW · 게재 시작 2026-09-22 · 캠페인 유형 검색(Search)"
ws["A2"].font = SUB
ws["H1"] = "캠페인 3개 게재중"; ws["H1"].font = GRN; ws["H1"].alignment = CEN
for i,w in enumerate([15,13,13,12,12,11,11,17],1):
    ws.column_dimensions[get_column_letter(i)].width = w

r = 4
ws.cell(row=r, column=1, value="광고 켜기 전 기준선").font = SEC
ws.cell(row=r, column=3, value="GA4 2026-09-09~09-20 평균 · 광고 시작 후 이 수준을 넘어야 증분 매출").font = SUB
r += 1
for i,(lab,val,fmt) in enumerate([("일 세션",round(avg_s),NUM),("일 거래수",round(avg_t,1),'#,##0.0'),
                                  ("일 매출",round(avg_r),WON),("객단가",round(avg_r/avg_t),WON)]):
    col = 1 + i*2
    c = ws.cell(row=r, column=col, value=lab); c.font = SUB
    c2 = ws.cell(row=r+1, column=col, value=val); c2.font = BIG; c2.number_format = fmt
    for cc in (c, c2): cc.fill = fill(CA_F)
    ws.cell(row=r, column=col+1).fill = fill(CA_F)
    ws.cell(row=r+1, column=col+1).fill = fill(CA_F)
ws.row_dimensions[r+1].height = 28

r += 4
ws.cell(row=r, column=1, value="캠페인").font = SEC
ws.cell(row=r, column=3, value="일 예산 합계 40,000원 · 월 약 120만원 · 청구 기준액 600,000원").font = SUB
r += 1
head(ws, r, ["캠페인","캠페인 ID","광고그룹","상태","일 예산","CPC 상한","키워드","문안"],
     [15,13,13,12,12,11,11,17])
r += 1
f0 = r
for n,cid,ag,b,ce,k in CAMPS:
    ws.cell(row=r, column=1, value=n).font = BD
    ws.cell(row=r, column=2, value=cid)
    ws.cell(row=r, column=3, value=ag)
    ws.cell(row=r, column=4, value="게재중").font = GRN
    ws.cell(row=r, column=4).alignment = CEN
    ws.cell(row=r, column=5, value=b).number_format = WON
    ws.cell(row=r, column=6, value=ce).number_format = WON
    ws.cell(row=r, column=7, value=k).alignment = CEN
    ws.cell(row=r, column=8, value="제목 15 · 설명 4")
    r += 1
ws.cell(row=r, column=1, value="합계").font = BD
ws.cell(row=r, column=5, value=f"=SUM(E{f0}:E{r-1})").number_format = WON
ws.cell(row=r, column=5).font = BD
ws.cell(row=r, column=7, value=f"=SUM(G{f0}:G{r-1})").font = BD
ws.cell(row=r, column=7).alignment = CEN
for col in range(1,9): ws.cell(row=r, column=col).fill = fill(BAND)

r += 2
ws.cell(row=r, column=1, value="공통 설정").font = SEC; r += 1
for k,v in [("위치","대한민국"),("언어","한국어 + 영어 (영문 브라우저 쓰는 국내 사용자 대응)"),
            ("입찰","클릭수 최대화 · 전환 30건 쌓이면 전환수 최대화로 변경")]:
    ws.cell(row=r, column=1, value=k).font = BD
    ws.cell(row=r, column=3, value=v); r += 1

r += 1
ws.cell(row=r, column=1, value="전환 추적").font = SEC; r += 1
for k,v in [("방식","GA4 전환 가져오기 · 카페24에 gtag 미설치"),
            ("전환 액션","젠제네틱스_속성 (web) purchase · 1개 · 기본 목표"),
            ("전환 값","GA4 실제 주문금액 → ROAS 계산 가능"),
            ("데이터 지연","24~48시간 · 비용은 있는데 전환 0인 날이 정상"),
            ("gclid 유실","없음 · 랜딩 3URL × 데스크톱/모바일 리다이렉트 0회 확인")]:
    ws.cell(row=r, column=1, value=k).font = BD
    ws.cell(row=r, column=3, value=v); r += 1

r += 1
ws.cell(row=r, column=1, value="지금 손봐야 할 것").font = SEC; r += 1
head(ws, r, ["등급","할 일","내용","","","","",""], [15,13,13,12,12,11,11,17])
r += 1
for sev,t,body in TODO:
    c = ws.cell(row=r, column=1, value=sev); c.alignment = CEN
    c.font = RED if sev=="위험" else (ORG if sev=="주의" else SUB)
    ws.cell(row=r, column=2, value=t).font = BD
    ws.cell(row=r, column=3, value=body).alignment = WRAP
    ws.row_dimensions[r].height = 30
    r += 1

# ═══ 2. 일별 기록 ═══
ws = wb.create_sheet("일별 기록")
ws.sheet_view.showGridLines = False
ws["A1"] = "일별 기록"; ws["A1"].font = TI
ws["A2"] = "노란 칸(노출·클릭·비용·전환·전환값)만 입력하세요. 회색 칸은 자동 계산됩니다."
ws["A2"].font = SUB
cols = ["날짜","요일","캠페인","노출","클릭","비용","전환","전환값","CTR","평균 CPC","CPA","ROAS","전환율","메모"]
head(ws, 4, cols, [12,7,13,10,9,12,9,13,9,11,11,9,9,26])
ws.freeze_panes = "D5"
r = 5
for d in range(DAYS):
    day = START + datetime.timedelta(days=d)
    for n,*_ in CAMPS:
        ws.cell(row=r, column=1, value=day.isoformat())
        ws.cell(row=r, column=2, value=DOW[day.weekday()]).alignment = CEN
        ws.cell(row=r, column=3, value=n)
        for col,fmt in ((4,NUM),(5,NUM),(6,WON),(7,'#,##0.0'),(8,WON)):
            c = ws.cell(row=r, column=col); c.fill = fill(IN_F); c.number_format = fmt
        for col,f,fmt in ((9,f'=IFERROR(E{r}/D{r},"")',PCT),(10,f'=IFERROR(F{r}/E{r},"")',WON),
                          (11,f'=IFERROR(F{r}/G{r},"")',WON),(12,f'=IFERROR(H{r}/F{r},"")','0.00"배"'),
                          (13,f'=IFERROR(G{r}/E{r},"")',PCT)):
            c = ws.cell(row=r, column=col, value=f); c.fill = fill(CA_F); c.number_format = fmt
        r += 1
tot = r
ws.cell(row=tot, column=1, value="누적").font = BD
for col,fmt in ((4,NUM),(5,NUM),(6,WON),(7,'#,##0.0'),(8,WON)):
    L = get_column_letter(col)
    c = ws.cell(row=tot, column=col, value=f"=SUM({L}5:{L}{tot-1})"); c.number_format = fmt; c.font = BD
for col,f,fmt in ((9,f'=IFERROR(E{tot}/D{tot},"")',PCT),(10,f'=IFERROR(F{tot}/E{tot},"")',WON),
                  (11,f'=IFERROR(F{tot}/G{tot},"")',WON),(12,f'=IFERROR(H{tot}/F{tot},"")','0.00"배"'),
                  (13,f'=IFERROR(G{tot}/E{tot},"")',PCT)):
    c = ws.cell(row=tot, column=col, value=f); c.number_format = fmt; c.font = BD
for col in range(1,15): ws.cell(row=tot, column=col).fill = fill(BAND)
r = tot + 2
ws.cell(row=r, column=1, value="지표 판단 기준").font = SEC; r += 1
for k,v in [("CTR","5% 이상 양호. 브랜드 키워드는 15%도 나옴"),
            ("평균 CPC","설정 상한 1,500 / 2,500 / 2,000원보다 낮아야 정상"),
            ("CPA","주문 1건 따오는 데 든 광고비"),
            ("ROAS","1배는 본전. 3배 이상이어야 남는 장사"),
            ("전환율","사이트 평균 1.2% 근처면 정상")]:
    ws.cell(row=r, column=1, value=k).font = BD
    ws.cell(row=r, column=3, value=v); r += 1

# ═══ 3. GA4 기준선 (차트 원본) ═══
ws = wb.create_sheet("GA4 기준선")
ws.sheet_view.showGridLines = False
ws["A1"] = "광고 켜기 전 매출 기준선"; ws["A1"].font = TI
ws["A2"] = "9월 9일부터 GA4 전자상거래가 켜져 그 전은 0으로 나옵니다. 실제로 안 팔린 게 아니라 측정이 없었던 것입니다."
ws["A2"].font = SUB
head(ws, 4, ["날짜","요일","세션","거래수","매출","객단가","전환율","비고"], [13,8,11,11,15,14,10,30])
ws.freeze_panes = "A5"
r = 5
for d,s,t,v in GA4:
    dt = datetime.date.fromisoformat(d)
    ws.cell(row=r, column=1, value=d[5:])
    ws.cell(row=r, column=2, value=DOW[dt.weekday()]).alignment = CEN
    ws.cell(row=r, column=3, value=s).number_format = NUM
    ws.cell(row=r, column=4, value=t).number_format = NUM
    ws.cell(row=r, column=5, value=v).number_format = WON
    ws.cell(row=r, column=6, value=f"=IFERROR(E{r}/D{r},0)").number_format = WON
    ws.cell(row=r, column=7, value=f"=IFERROR(D{r}/C{r},0)").number_format = PCT
    if d == "2026-09-21":
        ws.cell(row=r, column=8, value="이상치 — 세션도 최저(753)").font = ORG
    if d == "2026-09-22":
        ws.cell(row=r, column=8, value="광고 게재 시작일 / 집계중").font = SUB
    r += 1
last = 4 + len(GA4)
ws.cell(row=r, column=1, value="평균").font = BD
ws.cell(row=r, column=2, value="9/9~9/20").font = SUB
ws.cell(row=r, column=3, value=f"=ROUND(AVERAGE(C5:C{4+BASE_N}),0)").number_format = NUM
ws.cell(row=r, column=4, value=f"=ROUND(AVERAGE(D5:D{4+BASE_N}),1)").number_format = '#,##0.0'
ws.cell(row=r, column=5, value=f"=ROUND(AVERAGE(E5:E{4+BASE_N}),0)").number_format = WON
for col in range(1,9): ws.cell(row=r, column=col).fill = fill(BAND)

# ═══ 4. 키워드 ═══
ws = wb.create_sheet("키워드")
ws.sheet_view.showGridLines = False
ws["A1"] = "등록 키워드 42개"; ws["A1"].font = TI
ws["A2"] = "완전=이 검색어와 거의 똑같을 때만 / 구문=이 어구를 포함한 검색에 / 확장=관련 검색까지 폭넓게 (주의)"
ws["A2"].font = SUB
head(ws, 4, ["캠페인","매치타입","개수","키워드"], [15,12,8,96])
r = 5
for n,*_ in CAMPS:
    for mt in ("완전","구문","확장"):
        ks = [k for k,m in KW[n] if m == mt]
        if not ks: continue
        ws.cell(row=r, column=1, value=n).font = BD
        c = ws.cell(row=r, column=2, value=mt); c.alignment = CEN
        if mt == "확장": c.font = ORG
        ws.cell(row=r, column=3, value=len(ks)).alignment = CEN
        ws.cell(row=r, column=4, value=" · ".join(ks))
        r += 1
r += 1
ws.cell(row=r, column=1, value="주의").font = ORG
ws.cell(row=r, column=4, value="확장검색 3건은 예산이 새기 쉽습니다 → 검색어 보고서를 매일 보고 엉뚱한 검색은 제외 키워드로 막습니다.")
ws.cell(row=r+1, column=1, value="메모").font = BD
ws.cell(row=r+1, column=4, value="변형확장(+붓기 +영양제)은 구글이 2021년 폐지 → 확장검색으로 변환해 등록했습니다.")

# ═══ 5. 제외 키워드 ═══
ws = wb.create_sheet("제외 키워드")
ws.sheet_view.showGridLines = False
ws["A1"] = "제외 키워드 92개"; ws["A1"].font = TI
ws["A2"] = "캠페인 레벨 · 31단어 × 3캠페인 (브랜드 방어는 '무료'를 빼서 30개)"
ws["A2"].font = SUB
head(ws, 4, ["제외 사유","제외 키워드","브랜드 방어","붓기","칼륨"], [16,34,13,9,9])
r = 5
for grp, words in NEG:
    c = ws.cell(row=r, column=1, value=grp)
    c.font = RED if grp.startswith("의약품") else BD
    ws.cell(row=r, column=2, value=" · ".join(words))
    ws.cell(row=r, column=3, value="적용 (무료 제외)" if grp=="무료·체험" else "적용").alignment = CEN
    if grp == "무료·체험": ws.cell(row=r, column=3).font = ORG
    ws.cell(row=r, column=4, value="적용").alignment = CEN
    ws.cell(row=r, column=5, value="적용").alignment = CEN
    r += 1
r += 1
ws.cell(row=r, column=1, value="예외").font = ORG
ws.cell(row=r, column=2, value="브랜드 방어에서만 '무료'를 뺐습니다 — 「젠제네틱스 무료배송」은 구매 직전 검색이라 막으면 손해")

out = "/tmp/ads-dash.xlsx"
wb.save(out)
print("saved", out); print("tabs:", wb.sheetnames)
